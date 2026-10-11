"""Genera Plantilla_US_Abdomen_Renal_Pelvis_Prostata.docx con controles de contenido
(desplegables, texto y fecha) tipo "Programador" de Word."""
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from xml.sax.saxutils import escape
import itertools

AZUL = "1F3864"
_ids = itertools.count(1000)

def _rpr(bold=False, color=None):
    return ('<w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial" w:cs="Arial"/>'
            + ('<w:b/>' if bold else '') + (f'<w:color w:val="{color}"/>' if color else '')
            + '<w:sz w:val="22"/></w:rPr>')

def dropdown(p, titulo, opciones):
    items = "".join(f'<w:listItem w:displayText="{escape(o, {chr(34): "&quot;"})}" '
                    f'w:value="{escape(o, {chr(34): "&quot;"})}"/>' for o in opciones)
    xml = (f'<w:sdt {nsdecls("w")}><w:sdtPr>{_rpr()}<w:alias w:val="{escape(titulo)}"/>'
           f'<w:tag w:val="{escape(titulo)}"/><w:id w:val="{next(_ids)}"/>'
           f'<w:dropDownList>{items}</w:dropDownList></w:sdtPr>'
           f'<w:sdtContent><w:r>{_rpr()}<w:t xml:space="preserve">{escape(opciones[0])}</w:t></w:r>'
           f'</w:sdtContent></w:sdt>')
    p._p.append(parse_xml(xml))

def texto(p, titulo, ejemplo="Escriba aquí", bold=False):
    xml = (f'<w:sdt {nsdecls("w")}><w:sdtPr>{_rpr(bold)}<w:alias w:val="{escape(titulo)}"/>'
           f'<w:tag w:val="{escape(titulo)}"/><w:id w:val="{next(_ids)}"/><w:showingPlcHdr/><w:text/></w:sdtPr>'
           f'<w:sdtContent><w:r>{_rpr(bold, "808080")}<w:t xml:space="preserve">{escape(ejemplo)}</w:t></w:r>'
           f'</w:sdtContent></w:sdt>')
    p._p.append(parse_xml(xml))

def fecha(p):
    xml = (f'<w:sdt {nsdecls("w")}><w:sdtPr>{_rpr()}<w:alias w:val="Fecha"/><w:id w:val="{next(_ids)}"/>'
           '<w:showingPlcHdr/><w:date><w:dateFormat w:val="dd/MM/yyyy"/><w:lid w:val="es-NI"/>'
           '<w:calendar w:val="gregorian"/></w:date></w:sdtPr><w:sdtContent>'
           f'<w:r>{_rpr(color="808080")}<w:t>Seleccione fecha</w:t></w:r></w:sdtContent></w:sdt>')
    p._p.append(parse_xml(xml))

def run(p, t, bold=False, color=None, size=11):
    r = p.add_run(t); r.bold = bold; r.font.name = "Arial"; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    return r

def para(cont, space=2):
    p = cont.add_paragraph(); p.paragraph_format.space_after = Pt(space)
    p.paragraph_format.space_before = Pt(0); return p

def seccion(cont, t):
    p = para(cont, 3); p.paragraph_format.space_before = Pt(8)
    run(p, t, True, AZUL, 12)
    pPr = p._p.get_or_add_pPr()
    pPr.append(parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="8" w:space="1" w:color="{AZUL}"/></w:pBdr>'))

def organo(cont, nombre, opciones):
    p = para(cont); p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run(p, f"{nombre}: ", True, AZUL); dropdown(p, nombre, opciones)
    run(p, " "); texto(p, f"{nombre} (comentario adicional)", "[comentario opcional]")

def sombrear(celda, color):
    celda._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="{color}"/>'))

# ---------------------------------------------------------------- catálogo
NO_EVAL = "No evaluable adecuadamente por interposición de gas intestinal / ventana acústica limitada."
ORGANOS_ABD = [
 ("HÍGADO", ["De forma, contornos y dimensiones conservadas; ecotextura homogénea y ecogenicidad habitual, sin lesiones focales sólidas ni quísticas identificables. Bordes lisos.",
             "Hepatomegalia con ecotextura homogénea, sin lesiones focales identificables.",
             "Aumento difuso de la ecogenicidad parenquimatosa con preservación de la visualización de las paredes portales y del diafragma, compatible con esteatosis hepática grado I.",
             "Aumento difuso de la ecogenicidad con atenuación de la visualización de las paredes portales, compatible con esteatosis hepática grado II.",
             "Marcado aumento de la ecogenicidad con importante atenuación posterior que impide valorar diafragma y paredes portales, compatible con esteatosis hepática grado III.",
             "Contornos nodulares, ecotextura heterogénea y gruesa, con hipertrofia relativa del lóbulo caudado; hallazgos compatibles con hepatopatía crónica.",
             "Lesión focal descrita en comentario adicional.", NO_EVAL]),
 ("SISTEMA PORTAL", ["Vena porta de calibre conservado, permeable, con flujo hepatópeto.",
             "Vena porta de calibre aumentado, sugestivo de hipertensión portal.",
             "Vena porta con contenido ecogénico intraluminal sugestivo de trombosis.", NO_EVAL]),
 ("VENAS HEPÁTICAS", ["Permeables, de calibre habitual.", "Dilatadas, en relación con congestión hepática pasiva.", NO_EVAL]),
 ("VESÍCULA BILIAR", ["Adecuadamente distendida, de paredes finas y regulares, con contenido anecoico homogéneo; sin litiasis ni barro biliar. Murphy ecográfico negativo.",
             "Distendida, con imagen(es) ecogénica(s) intraluminal(es) móvil(es) con sombra acústica posterior, compatibles con colelitiasis; paredes finas. Murphy negativo.",
             "Distendida, con litiasis, engrosamiento parietal difuso, edema mural y Murphy ecográfico positivo, compatible con colecistitis aguda litiásica.",
             "Contenido ecogénico declive sin sombra acústica, compatible con barro biliar.",
             "Imagen(es) ecogénica(s) adherida(s) a la pared, fija(s), sin sombra acústica, compatible(s) con pólipo(s) vesicular(es).",
             "Contraída (paciente posprandial), lo que limita su evaluación.",
             "No visualizada; antecedente de colecistectomía.", NO_EVAL]),
 ("VÍAS BILIARES", ["Vías biliares intra y extrahepáticas sin dilatación. Colédoco de calibre conservado.",
             "Dilatación de las vías biliares intrahepáticas y del colédoco.",
             "Colédoco dilatado con imagen ecogénica intraluminal con sombra acústica, sugestiva de coledocolitiasis.", NO_EVAL]),
 ("PÁNCREAS", ["Cabeza, cuerpo y cola de morfología y ecogenicidad conservadas; conducto pancreático principal no dilatado, sin lesiones focales identificables.",
             "Aumentado de tamaño, hipoecoico, con borramiento de planos grasos peripancreáticos, sugestivo de pancreatitis aguda.",
             "Parcialmente evaluado por interposición de gas; segmentos visibles sin alteraciones.", NO_EVAL]),
 ("BAZO", ["De dimensiones y ecotextura conservadas, sin lesiones focales identificables.",
             "Esplenomegalia homogénea, sin lesiones focales.", "Bazo accesorio (esplenúnculo) en hilio esplénico, hallazgo incidental.",
             "No visualizado; antecedente de esplenectomía."]),
 ("RETROPERITONEO Y GRANDES VASOS", ["Aorta abdominal de calibre conservado, sin dilataciones aneurismáticas; vena cava inferior permeable. Sin adenomegalias retroperitoneales evaluables.",
             "Aorta abdominal con placas ateromatosas parietales calcificadas, sin dilatación aneurismática.",
             "Dilatación aneurismática de la aorta abdominal (ver medidas).", NO_EVAL]),
 ("CAVIDAD PERITONEAL", ["Sin evidencia de líquido libre ni colecciones.", "Líquido libre en escasa cuantía.",
             "Líquido libre en moderada cuantía.", "Ascitis abundante.", "Colección organizada (ver comentario y medidas)."]),
]
ORGANOS_RENAL = [
 ("RIÑÓN DERECHO", ["De situación, forma, tamaño y contornos conservados; adecuada diferenciación corticomedular, grosor parenquimatoso preservado; sistema pielocalicial sin dilatación; sin litiasis ni lesiones focales.",
             "Imagen ecogénica con sombra acústica posterior en grupo calicial, compatible con litiasis renal no obstructiva.",
             "Ectasia pielocalicial leve (hidronefrosis grado I).", "Hidronefrosis moderada (grado II).", "Hidronefrosis severa (grado III–IV) con adelgazamiento cortical.",
             "Quiste cortical simple (Bosniak I).", "Aumento difuso de la ecogenicidad cortical con pérdida de la diferenciación corticomedular, compatible con nefropatía médica crónica.",
             "Disminuido de tamaño con adelgazamiento cortical difuso.", "No visualizado (agenesia / nefrectomía)."]),
 ("RIÑÓN IZQUIERDO", None),
 ("URÉTERES", ["Porciones proximales y distales evaluables no dilatadas.", "Uréter derecho dilatado.", "Uréter izquierdo dilatado.",
             "Litiasis en unión ureterovesical (ver comentario).", NO_EVAL]),
 ("VEJIGA URINARIA", ["Adecuadamente distendida, paredes finas y regulares, contenido anecoico; sin litiasis, masas ni divertículos.",
             "Engrosamiento parietal difuso con trabeculación, sugestivo de vejiga de esfuerzo.",
             "Contenido ecogénico en suspensión (sedimento).", "Imagen ecogénica móvil con sombra acústica, compatible con litiasis vesical.",
             "Lesión polipoide parietal (ver comentario).", "Insuficientemente distendida, lo que limita su evaluación.", "Sonda vesical in situ."]),
 ("RESIDUO POSMICCIONAL", ["No significativo.", "Significativo (ver volumen en cuadro de mediciones).", "No evaluado."]),
]
ORGANOS_RENAL[1] = ("RIÑÓN IZQUIERDO", ORGANOS_RENAL[0][1])
ORGANOS_FEM = [
 ("ÚTERO", ["En anteversoflexión, de forma, contornos y dimensiones conservadas; miometrio de ecotextura homogénea, sin lesiones focales.",
             "En retroversoflexión, de forma y dimensiones conservadas, miometrio homogéneo.",
             "Aumentado de tamaño, con imagen(es) nodular(es) hipoecoica(s) bien definida(s), compatibles con miomatosis uterina (FIGO en comentario).",
             "Miometrio heterogéneo, globuloso, con asimetría de paredes, sugestivo de adenomiosis.",
             "Útero de características posmenopáusicas, atrófico.", "No visualizado; antecedente de histerectomía."]),
 ("ENDOMETRIO", ["Lineal, ecogénico, de grosor acorde a la fase del ciclo menstrual.", "Trilaminar (fase proliferativa).", "Ecogénico homogéneo (fase secretora).",
             "Engrosado para la edad / estado hormonal.", "Contenido líquido en cavidad endometrial.", "Dispositivo intrauterino normoinserto.", "Dispositivo intrauterino descendido."]),
 ("OVARIO DERECHO", ["De tamaño y morfología conservados con folículos de aspecto fisiológico (O-RADS 1).",
             "Imagen quística unilocular anecoica de paredes finas (quiste simple, O-RADS 2).",
             "Quiste con ecos internos en vidrio esmerilado, sugestivo de endometrioma (O-RADS 2).",
             "Quiste con patrón reticular, sugestivo de quiste hemorrágico (O-RADS 2).",
             "Aumento de volumen con múltiples folículos periféricos, morfología poliquística.",
             "No visualizado (posmenopausia / gas intestinal)."]),
 ("OVARIO IZQUIERDO", None),
 ("ANEXOS Y FONDO DE SACO", ["Regiones anexiales sin masas; fondo de saco de Douglas libre.", "Escaso líquido libre en fondo de saco de Douglas, probablemente fisiológico.",
             "Líquido libre en moderada cantidad en fondo de saco.", "Masa anexial (ver comentario)."]),
]
ORGANOS_FEM[3] = ("OVARIO IZQUIERDO", ORGANOS_FEM[2][1])
ORGANOS_MASC = [
 ("PRÓSTATA", ["De forma, contornos y dimensiones conservadas, ecotextura homogénea, sin lesiones focales; peso estimado dentro de parámetros esperados.",
             "Aumentada de tamaño, con ecotextura heterogénea y calcificaciones en zona de transición, compatible con hiperplasia prostática.",
             "Aumentada de tamaño con protrusión del lóbulo medio hacia la base vesical.",
             "Calcificaciones intraprostáticas (prostatolitos).", "Lesión focal hipoecoica periférica (ver comentario).", "No evaluada."]),
 ("VESÍCULAS SEMINALES", ["Simétricas, de aspecto habitual.", "Asimétricas / dilatadas.", "No evaluadas."]),
]

MEDIDAS = [("Hígado (lóbulo derecho)", "cm", False), ("Vena porta", "mm", False), ("Vesícula biliar (L × AP × T)", "cm", True),
           ("Pared vesicular", "mm", False), ("Colédoco", "mm", False), ("Bazo (L × AP × T)", "cm", True),
           ("Riñón derecho (L × AP × T)", "cm", True), ("Parénquima renal derecho", "mm", False),
           ("Riñón izquierdo (L × AP × T)", "cm", True), ("Parénquima renal izquierdo", "mm", False),
           ("Aorta abdominal", "mm", False), ("Vejiga prevaciado (L × AP × T)", "cm", True),
           ("Vejiga posvaciado (L × AP × T)", "cm", True), ("Útero (L × AP × T)", "cm", True), ("Endometrio", "mm", False),
           ("Ovario derecho (L × AP × T)", "cm", True), ("Ovario izquierdo (L × AP × T)", "cm", True),
           ("Próstata (L × AP × T)", "cm", True), ("Lesión / quiste / colección (L × AP × T)", "cm", True)]

# ---------------------------------------------------------------- documento
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.59), Cm(27.94)
sec.left_margin = sec.right_margin = Cm(1.6); sec.top_margin = Cm(1.4); sec.bottom_margin = Cm(2.3)
sec.different_first_page_header_footer = True
# encabezado compacto páginas 2+
hp = sec.header.paragraphs[0]
run(hp, "HOSPITAL REGIONAL ESCUELA CAMILO ORTEGA SAAVEDRA · Radiología e Imágenes Médicas · Juigalpa, Chontales", True, AZUL, 8)
for f in (sec.footer, sec.first_page_footer):
    fp = f.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(fp, "Dr. Néstor José Mejía González · Especialista en Radiología e Imagen · Código MINSA 37294", False, "595959", 8)

tabla = doc.add_table(rows=1, cols=2); tabla.autofit = False
izq, der = tabla.rows[0].cells
izq.width, der.width = Cm(4.4), Cm(14.0)
for c, w in ((izq, 4.4), (der, 14.0)):
    c._tc.get_or_add_tcPr().append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{int(w*567)}" w:type="dxa"/>'))
sombrear(izq, "E8EEF7")

# columna lateral
p = izq.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "[LOGO HOSPITAL]\n", True, AZUL, 9); run(p, "\n[SELLO NM RADIOLOGÍA]", True, AZUL, 9)
p = para(izq, 6); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "Reemplace con: clic derecho › Cambiar imagen", False, "808080", 7)
def dato(etq, fn):
    p = para(izq, 0); run(p, etq, True, AZUL, 9)
    p = para(izq, 5); fn(p)
dato("PACIENTE:", lambda p: texto(p, "Paciente", "Nombre y apellidos", True))
dato("EDAD:", lambda p: texto(p, "Edad", "__ años"))
dato("SEXO:", lambda p: dropdown(p, "Sexo", ["Femenino", "Masculino"]))
dato("MÉDICO SOLICITANTE:", lambda p: texto(p, "Médico solicitante", "Nombre del médico"))
dato("PROCEDENCIA:", lambda p: dropdown(p, "Procedencia", ["Consulta externa", "Emergencia", "Hospitalización", "Medicina interna", "Cirugía", "Ginecoobstetricia", "Pediatría", "Urología", "UCI", "Externo"]))
dato("MODALIDAD:", lambda p: dropdown(p, "Modalidad", ["Ultrasonido", "Ultrasonido Doppler"]))
dato("FECHA:", fecha)
dato("HORA:", lambda p: texto(p, "Hora", "hh:mm"))
p = para(izq, 0); p.paragraph_format.space_before = Pt(18)
run(p, "Dr. Néstor José Mejía González\n", True, AZUL, 8); run(p, "Especialista en Radiología e Imagen\nCódigo MINSA 37294", False, "404040", 8)

# área clínica derecha
p = der.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "INFORME RADIOLÓGICO", True, AZUL, 16)
p = para(der, 6); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "Departamento de Radiología e Imágenes Médicas · Juigalpa, Chontales", False, "404040", 9)
seccion(der, "ESTUDIO / INDICACIÓN CLÍNICA")
p = para(der); run(p, "Estudio: ", True)
dropdown(p, "Estudio", ["Ultrasonido abdominal completo", "Ultrasonido abdominal superior", "Ultrasonido renal y de vías urinarias",
                        "Ultrasonido general femenino (abdomen y pelvis)", "Ultrasonido general masculino (abdomen y próstata)",
                        "Ultrasonido pélvico transabdominal", "Ultrasonido prostático transabdominal", "Ultrasonido vesicoprostático"])
p = para(der); run(p, "Indicación clínica: ", True); texto(p, "Indicación clínica", "Motivo del estudio")
p = para(der); run(p, "Técnica: ", True)
dropdown(p, "Técnica", ["Estudio realizado en tiempo real con transductor convexo multifrecuencia, en cortes longitudinales, transversales y oblicuos.",
                        "Estudio realizado con transductor convexo y endocavitario.", "Estudio realizado con vejiga en repleción y posvaciado."])
seccion(der, "HALLAZGOS")
p = para(der); run(p, "Borre los bloques que no correspondan al estudio (clic en la etiqueta del control › Suprimir).", False, "808080", 8)
for n, o in ORGANOS_ABD: organo(der, n, o)

# continuación a ancho completo
for titulo, lista in (("RIÑONES Y VÍAS URINARIAS", ORGANOS_RENAL), ("PELVIS FEMENINA", ORGANOS_FEM), ("PRÓSTATA / PELVIS MASCULINA", ORGANOS_MASC)):
    seccion(doc, titulo)
    for n, o in lista: organo(doc, n, o)

seccion(doc, "CUADRO DE MEDICIONES")
t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
for c, h in zip(t.rows[0].cells, ("Estructura", "Medidas", "Unidad", "Volumen (cm³)")):
    c.paragraphs[0].text = ""; run(c.paragraphs[0], h, True, "FFFFFF", 10); sombrear(c, AZUL)
for i, (n, u, vol) in enumerate(MEDIDAS):
    r = t.add_row().cells
    run(r[0].paragraphs[0], n, False, None, 10)
    texto(r[1].paragraphs[0], f"Medida {n}", "__ × __ × __" if vol else "__")
    run(r[2].paragraphs[0], u, False, None, 10)
    if vol: texto(r[3].paragraphs[0], f"Volumen {n}", "L×AP×T×0.523")
    else: run(r[3].paragraphs[0], "—", False, "808080", 10)
    if i % 2: [sombrear(c, "F2F5FA") for c in r]
for row in t.rows:
    for c, w in zip(row.cells, (7.0, 5.0, 2.0, 4.4)):
        c._tc.get_or_add_tcPr().append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{int(w*567)}" w:type="dxa"/>'))
t.autofit = False

seccion(doc, "CONCLUSIONES")
for i in range(1, 5):
    p = para(doc); run(p, f"{i}. ", True)
    if i == 1:
        dropdown(p, "Conclusión principal", ["Ultrasonido abdominal sin alteraciones estructurales demostrables.",
            "Esteatosis hepática.", "Colelitiasis sin signos de colecistitis aguda.", "Colecistitis aguda litiásica.",
            "Hepatopatía crónica.", "Nefrolitiasis.", "Hidronefrosis.", "Nefropatía médica crónica.",
            "Miomatosis uterina.", "Hiperplasia prostática.", "Ultrasonido renal y vesical sin alteraciones demostrables.",
            "Ultrasonido pélvico sin alteraciones demostrables.", "Ver texto libre."])
        run(p, " "); texto(p, "Conclusión 1 detalle", "[detalle]")
    else:
        texto(p, f"Conclusión {i}", "[conclusión adicional / borrar si no aplica]")
seccion(doc, "ANÁLISIS ETIOLÓGICO / CAUSA PROBABLE")
p = para(doc); texto(p, "Análisis etiológico", "Relación breve del problema principal con sus causas radiológicamente plausibles.")

p = para(doc, 0); p.paragraph_format.space_before = Pt(36); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(p, "_______________________________\n"); run(p, "Dr. Néstor José Mejía González\n", True, AZUL)
run(p, "Médico Especialista en Radiología e Imágenes\nCódigo MINSA: 37294")

import os; doc.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "Plantilla_US_Abdomen_Renal_Pelvis_Prostata.docx"))
print("ok")
